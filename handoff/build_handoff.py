"""Reproducible Word handoff builder plus source/asset inventories. No secret files read.
Run from repository root: python handoff/build_handoff.py
Requires python-docx, Pillow. LibreOffice and PyMuPDF are used by validate_handoff.py.
"""
from pathlib import Path
from hashlib import sha256
import csv
import json
import re
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'handoff'
SOURCE = OUT / 'RUMI_DEVELOPER_HANDOFF.md'
DOCX = OUT / 'Rumi_Developer_Handoff.docx'
NATIVE = ROOT / 'ios/Nudge'
WEB = ROOT / 'web/src'
# Only allowlisted source extensions/roots; never Config or environment bindings.
sources = [p for root in [NATIVE, WEB, ROOT/'functions'] for p in root.rglob('*') if p.suffix in ['.swift','.ts','.tsx','.css','.metal'] and 'node_modules' not in p.parts and p.name not in ['Config.swift','Config.ts']]
contents = {str(p.relative_to(ROOT)):p.read_text(errors='replace') for p in sources}
assets = []
for root in [NATIVE/'Assets.xcassets', NATIVE/'Resources', NATIVE/'Fonts', ROOT/'web/public', ROOT/'tmp/brand-delivery']:
    for path in sorted(root.rglob('*')):
        if path.suffix.lower() not in ['.png','.jpg','.jpeg','.svg','.mp4','.m4a','.mp3','.ttf','.otf'] or not path.is_file(): continue
        raw = path.read_bytes()
        name = path.stem
        refs = [p for p,c in contents.items() if name in c]
        info = {'path':str(path.relative_to(ROOT)), 'bytes':len(raw), 'sha256':sha256(raw).hexdigest(), 'reference_status':'source token matches' if refs else 'no static token match; dynamic/persisted references may exist', 'consumers':'; '.join(refs), 'dimensions':'', 'provenance':'Not established by packaging; see handoff family notes'}
        if path.suffix.lower() in ['.png','.jpg','.jpeg']:
            with Image.open(path) as image: info['dimensions'] = f'{image.width}x{image.height} {image.mode}'
        assets.append(info)
with (OUT/'asset_inventory.csv').open('w',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=list(assets[0]),lineterminator='\n'); writer.writeheader(); writer.writerows(assets)
(OUT/'source_inventory.json').write_text(json.dumps([{'path':p,'sha256':sha256(c.encode()).hexdigest()} for p,c in sorted(contents.items())],indent=2))

text=SOURCE.read_text()
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
normal=doc.styles['Normal']; normal.font.name='Liberation Sans'; normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(7)
for level in range(1,4):
    style=doc.styles[f'Heading {level}']; style.font.name='Liberation Sans'; style.font.color.rgb=RGBColor.from_string('164A3D'); style.font.size=Pt(24 if level==1 else 16 if level==2 else 12); style.paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0]; header.text='RUMI  /  DEVELOPER HANDOFF'; header.style=doc.styles['Caption']; header.runs[0].font.color.rgb=RGBColor.from_string('526D61')
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Synthetic demo · Production gaps explicit   |   ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def inline(paragraph,line):
    # Keep paths literal; implement clickable links, inline code and bold without raw Markdown.
    pattern=r'(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for part in re.split(pattern,line):
        link=re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',part)
        if link:
            label,target=link.groups()
            if not target.startswith(('http://','https://','#')): target='../'+target
            relation=paragraph.part.relate_to(target,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            node=OxmlElement('w:hyperlink'); node.set(qn('r:id'),relation)
            run=OxmlElement('w:r'); props=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'287D66'); props.append(color); run.append(props)
            t=OxmlElement('w:t'); t.text=label; run.append(t); node.append(run); paragraph._p.append(node)
        elif part.startswith('**'):
            paragraph.add_run(part[2:-2]).bold=True
        elif part.startswith('`'):
            run=paragraph.add_run(part[1:-1]); run.font.name='Liberation Mono'; run.font.size=Pt(8)
        else: paragraph.add_run(part)

def table(lines):
    rows=[[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[-: ]+',c or '-') for c in r)]
    t=doc.add_table(rows=1, cols=len(rows[0])); t.style='Light Shading Accent 1'; t.autofit=False
    for i,value in enumerate(rows[0]): t.rows[0].cells[i].text=value
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows[1:]:
        cells=t.add_row().cells
        for i,value in enumerate(row[:len(cells)]): inline(cells[i].paragraphs[0],value)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                for r in p.runs:r.font.size=Pt(8)
    doc.add_paragraph()

lines=text.splitlines(); index=0; in_code=False
while index<len(lines):
    line=lines[index]
    if line.startswith('```'): in_code=not in_code; index+=1; continue
    if in_code:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); run=p.add_run(line); run.font.name='Liberation Mono'; run.font.size=Pt(8)
    elif line=='<!-- pagebreak -->': doc.add_page_break()
    elif line=='<!-- screenshots -->':
        ledger=json.loads((ROOT/'screenshots/privia-web/manifest.json').read_text())
        extra=OUT/'live_ai_verification.json'
        if extra.exists(): ledger['screens'].extend(json.loads(extra.read_text()).get('screens',[]))
        for i,screen in enumerate(ledger['screens']):
            doc.add_heading(f"{i+1:02d} · {screen['file'].removesuffix('.png')}",2)
            p=doc.add_paragraph(screen['caption']); p.paragraph_format.keep_with_next=True
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT/'screenshots/privia-web'/screen['file']),width=Inches(3.25))
            doc.add_paragraph('Actual Chromium web capture · Elena, synthetic · sample actions only',style='Caption')
            if i<len(ledger['screens'])-1: doc.add_page_break()
    elif line=='<!-- asset-register -->':
        doc.add_paragraph('The companion asset_inventory.csv provides all paths, sizes, SHA-256 hashes, dimensions and static source-token consumers. The compact register below lists every packaged asset; token matches are not proof that a currently reachable screen uses it.')
        for root in ['ios/Nudge/Assets.xcassets','ios/Nudge/Resources','ios/Nudge/Fonts','web/public','tmp/brand-delivery']:
            doc.add_heading(root,3)
            selected=[a for a in assets if a['path'].startswith(root+'/')]
            for a in selected:
                short=a['path'][len(root)+1:]
                p=doc.add_paragraph(); inline(p,f"`{short}` — {a['bytes']:,} bytes; {a['dimensions'] or 'media/font'}; {'source references' if a['consumers'] else 'no static reference found'}.")
    elif line.startswith('|'):
        group=[]
        while index<len(lines) and lines[index].startswith('|'): group.append(lines[index]);index+=1
        table(group);continue
    elif line.startswith('#'):
        level=len(line)-len(line.lstrip('#')); title=line[level:].strip()
        doc.add_heading(title,min(level,3))
    elif line.startswith('- '): inline(doc.add_paragraph(style='List Bullet'),line[2:])
    elif re.match(r'^\d+\. ',line): inline(doc.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',line))
    elif line.strip(): inline(doc.add_paragraph(),line)
    index+=1

doc.core_properties.title='Rumi — Developer Handoff'
doc.core_properties.subject='Web Privia demo, native iOS, services, design/assets and production release gaps'
doc.core_properties.author='Rumi project'
doc.core_properties.comments='Source-backed handoff, not a QMS record or production certification.'
doc.save(DOCX)
print(json.dumps({'docx':str(DOCX.relative_to(ROOT)),'assetRows':len(assets),'sourceRows':len(contents),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'bytes':DOCX.stat().st_size}))
